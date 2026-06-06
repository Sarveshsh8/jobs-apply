from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .extractor import JobListing

HEADERS = [
    "Title", "Company", "Location", "Location Verdict",
    "Location Reason", "Posted Date", "Link", "Job Description",
]


def save_to_excel(
    results: list[tuple[JobListing, str, str]],
    output_dir: Path = Path("output"),
) -> Path | None:
    if not results:
        print("\nNo new matching jobs found.")
        return None

    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M")
    path = output_dir / f"jobs_{timestamp}.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    wrap = Alignment(wrap_text=True, vertical="top")

    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = wrap

    for row_idx, (job, verdict, reason) in enumerate(results, 2):
        ws.cell(row=row_idx, column=1, value=job.title).alignment = wrap
        ws.cell(row=row_idx, column=2, value=job.company).alignment = wrap
        ws.cell(row=row_idx, column=3, value=job.location).alignment = wrap
        ws.cell(row=row_idx, column=4, value=verdict).alignment = wrap
        ws.cell(row=row_idx, column=5, value=reason).alignment = wrap
        ws.cell(row=row_idx, column=6, value=job.posted_date).alignment = wrap
        ws.cell(row=row_idx, column=7, value=job.link).alignment = wrap
        ws.cell(row=row_idx, column=8, value=job.description[:30000]).alignment = wrap

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 25
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 45
    ws.column_dimensions["H"].width = 60

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    wb.save(path)
    print(f"\n{'='*60}")
    print(f" Found {len(results)} jobs → saved to {path}")
    print(f"{'='*60}")
    return path
